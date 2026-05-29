import io
import math
import re
from typing import Any
from openpyxl import load_workbook

CHUNK_SIZE = 1000


def _clean_value(val: Any, is_int: bool = False) -> Any:
    if val is None:
        return None
    if isinstance(val, float) and math.isnan(val):
        return None
    s = str(val).strip()
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.IGNORECASE)
    s = s.strip()
    if s == "":
        return None
    # N/A, n/a, N/D, etc. se conservan tal cual (no se convierten a NULL)
    if is_int:
        try:
            return int(float(s))
        except (ValueError, TypeError):
            return None
    return s


def validate_excel_columns(file_bytes: bytes, col_map: dict) -> dict:
    """
    Detecta columnas faltantes y registros con campos vacíos.
    Retorna:
    - columnas_faltantes: campos cuyo índice supera el ancho del archivo
    - registros_con_vacios: [ { fila, campos_vacios: [col, ...], codigo_referencia } ]
    - total_filas: filas de datos (sin encabezado, sin filas completamente vacías)
    """
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    worksheet = workbook.active
    total_cols = worksheet.max_column or 0

    missing = [db_col for db_col, idx in col_map.items() if int(idx) >= total_cols]
    available = {db_col: int(idx) for db_col, idx in col_map.items() if int(idx) < total_cols}

    # índice de codigo_referencia para mostrarlo como identificador en la UI
    cod_ref_idx = available.get("codigo_referencia")

    registros_con_vacios = []
    data_rows = 0

    for row_idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            continue  # saltar encabezado
        if all(v is None for v in row):
            continue
        data_rows += 1
        campos_vacios = [
            db_col
            for db_col, col_idx in available.items()
            if (row[col_idx] if col_idx < len(row) else None) in (None, "")
            or str(row[col_idx] if col_idx < len(row) else "").strip() == ""
        ]
        if campos_vacios:
            cod_ref = None
            if cod_ref_idx is not None and cod_ref_idx < len(row):
                cod_ref = row[cod_ref_idx]
            registros_con_vacios.append({
                "fila": row_idx,
                "campos_vacios": campos_vacios,
                "codigo_referencia": str(cod_ref) if cod_ref is not None else None,
            })

    workbook.close()
    return {
        "columnas_faltantes": missing,
        "registros_con_vacios": registros_con_vacios,
        "total_filas": data_rows,
    }


def stream_excel_chunks(file_bytes: bytes, col_map: dict, date_cols: list):
    workbook = load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    total_cols = worksheet.max_column or 0
    missing = [db_col for db_col, idx in col_map.items() if int(idx) >= total_cols]
    available = {db_col: int(idx) for db_col, idx in col_map.items() if int(idx) < total_cols}

    rows_total = worksheet.max_row or 0

    def chunks():
        chunk = []
        try:
            for excel_row, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                record = {
                    db_col: _clean_value(
                        row[col_idx] if col_idx < len(row) else None,
                        is_int=(db_col in date_cols),
                    )
                    for db_col, col_idx in available.items()
                }
                if all(value is None for value in record.values()):
                    continue
                record["_excel_row"] = excel_row
                chunk.append(record)
                if len(chunk) >= CHUNK_SIZE:
                    yield chunk
                    chunk = []

            if chunk:
                yield chunk
        finally:
            workbook.close()

    return chunks(), rows_total, missing


def read_excel_chunks(file_bytes: bytes, col_map: dict, date_cols: list):
    chunk_iter, rows_total, missing = stream_excel_chunks(file_bytes, col_map, date_cols)
    chunks = []
    for chunk in chunk_iter:
        chunks.append(chunk)
    return chunks, rows_total, missing
